"""Parse the World Bank's historical income classification workbook (OGHIST.xlsx).

Sheet "Country Analytical History" has a header row of data (calendar) years, a row of the
World Bank fiscal years the classification applied to, then one row per economy with codes
L / LM / UM / H. ".." means not available and an empty cell means not classified (e.g.
Venezuela since 2021); both become "INX" (not classified), with the original code kept.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl

CODES = {"L": "LIC", "LM": "LMC", "LM*": "LMC", "UM": "UMC", "H": "HIC"}
NOT_CLASSIFIED = "INX"


class IncomeHistoryError(ValueError):
    """The workbook does not have the expected layout."""


def parse_oghist(body: bytes, sheet: str) -> list[dict[str, Any]]:
    """One record per economy x data year: iso3, data_year, fiscal_year, code, income_group_id."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:  # zip/xml errors: not a workbook at all
        raise IncomeHistoryError(f"not a readable .xlsx workbook: {exc}") from exc
    if sheet not in wb.sheetnames:
        raise IncomeHistoryError(f"sheet {sheet!r} not found (have {wb.sheetnames})")
    rows = list(wb[sheet].iter_rows(values_only=True))

    def find(label: str) -> tuple[Any, ...]:
        for row in rows:
            if len(row) > 1 and isinstance(row[1], str) and row[1].strip().startswith(label):
                return row
        raise IncomeHistoryError(f"no {label!r} header row in sheet {sheet!r}")

    years_row, fiscal_row = find("Data for calendar year"), find("Bank's fiscal year")
    columns = [i for i, v in enumerate(years_row) if isinstance(v, int | float)]
    if not columns:
        raise IncomeHistoryError("no year columns found")

    records = []
    for row in rows:
        iso3 = row[0].strip() if isinstance(row[0], str) else ""
        if len(iso3) != 3 or not iso3.isalpha() or not iso3.isupper():
            continue  # titles, threshold rows, footnotes, former economies like "YUGf"
        for i in columns:
            raw = row[i] if i < len(row) else None
            code = str(raw).strip() if raw is not None else ""
            records.append({
                "country_iso3": iso3,
                "data_year": int(years_row[i]),
                "fiscal_year": str(fiscal_row[i]).strip() if fiscal_row[i] else None,
                "source_code": code,
                "income_group_id": CODES.get(code, NOT_CLASSIFIED),
            })  # fmt: skip
    if not records:
        raise IncomeHistoryError("no economy rows found")
    return records
